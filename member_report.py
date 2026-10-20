#!/usr/bin/env python3
"""
Steel Member Survey - PC tool
=============================

Merges the project zips exported by the iPad "Steel Member Survey" app and builds the
Word report (same layout the iPad produces).

    pip install python-docx pillow

    # merge any number of exports into one project folder (duplicates by area+part are
    # resolved automatically: the first copy is kept unless a later one is more complete)
    python member_report.py --merge Merged export1.zip export2.zip ...

    # build the Word report from a merged (or single, unzipped) project
    python member_report.py --build Merged/project.json Report.docx

    # or both at once
    python member_report.py --merge Merged a.zip b.zip --build Report.docx
"""
import json, os, re, shutil, sys, tempfile, zipfile

from PIL import Image, ImageOps


# ----------------------------------------------------------------------------- loading
def load_project(path):
    """Open project.json, or an iPad zip (unpacked into <zip folder>/_unpacked/<zip name>/)."""
    if path.lower().endswith(".zip"):
        dest = os.path.join(os.path.dirname(os.path.abspath(path)), "_unpacked", os.path.splitext(os.path.basename(path))[0])
        os.makedirs(dest, exist_ok=True)
        with zipfile.ZipFile(path) as z:
            z.extractall(dest)
        cand = [os.path.join(r, "project.json") for r, _, fs in os.walk(dest) if "project.json" in fs]
        if not cand:
            raise ValueError(f"project.json not found in {path}")
        path = sorted(cand, key=len)[0]
    with open(path, encoding="utf-8") as f:
        p = json.load(f)
    if p.get("kind") != "member":
        raise ValueError(f"{path} is not a Steel Member Survey export (kind={p.get('kind')!r})")
    base = os.path.dirname(os.path.abspath(path))
    p["_photo_dir"] = os.path.join(base, p.get("photo_dir") or "photos")
    p["_path"] = path
    for m in p["modillions"]:
        m.setdefault("overall", []); m.setdefault("surfaces", {}); m.setdefault("checked", list(m["surfaces"].keys()))
    return p


def label(m, abbr="A"):
    return f"{abbr}{m['zone']}-{m['number']}"


def sort_key(m):
    def k(v):
        v = str(v).strip()
        return (0, int(v), "") if v.isdigit() else (1, 0, v.lower())
    return (k(m["zone"]), k(m["number"]))


def photo_ids(m):
    ids = list(m.get("overall", []))
    for sf in m.get("surfaces", {}).values():
        if sf.get("photo"): ids.append(sf["photo"])
        if sf.get("gauge"): ids.append(sf["gauge"])
    return [i for i in ids if i]


def completeness(m):
    sfs = m.get("surfaces", {})
    return (len(photo_ids(m)), sum(len((sf.get("readings") or "").strip()) + len((sf.get("notes") or "").strip()) for sf in sfs.values()) + len((m.get("notes") or "").strip()), len(sfs))


# ----------------------------------------------------------------------------- merging
def merge(out_dir, inputs):
    photos = os.path.join(out_dir, "photos"); os.makedirs(photos, exist_ok=True)
    merged = None; existing = set(); by_label = {}
    for src_path in inputs:
        src = load_project(src_path)
        if merged is None:
            merged = {k: v for k, v in src.items() if not k.startswith("_")}
            merged["modillions"] = []; merged["zones"] = {}; merged["photo_dir"] = "photos"
        abbr = (merged["settings"].get("group_name") or "Area")[0].upper()

        def copy_in(rel):
            if not rel: return ""
            s = os.path.join(src["_photo_dir"], rel)
            if not os.path.isfile(s): return ""
            base, ext = os.path.splitext(os.path.basename(s)); name = base + ext; k = 1
            while name.lower() in existing:
                k += 1; name = f"{base}_{k}{ext}"
            existing.add(name.lower()); shutil.copy2(s, os.path.join(photos, name)); return name

        def bring_in(m):
            m = dict(m); m["overall"] = [copy_in(p) for p in m.get("overall", []) if p]
            m["surfaces"] = {k: dict(sf, photo=copy_in(sf.get("photo")), gauge=copy_in(sf.get("gauge"))) for k, sf in m.get("surfaces", {}).items()}
            return m

        for z, info in src.get("zones", {}).items():
            cur = merged["zones"].get(z)
            if not cur or (not cur.get("description") and not cur.get("image")):
                info = dict(info); info["image"] = copy_in(info.get("image", "")); merged["zones"][z] = info
        c = {"added": 0, "skipped": 0, "replaced": 0}
        for m in src["modillions"]:
            l = label(m, abbr)
            if l in by_label:
                ours = merged["modillions"][by_label[l]]
                if completeness(m) > completeness(ours):
                    merged["modillions"][by_label[l]] = bring_in(m); c["replaced"] += 1
                else:
                    c["skipped"] += 1
                continue
            merged["modillions"].append(bring_in(m)); by_label[l] = len(merged["modillions"]) - 1; c["added"] += 1
        print(f"{os.path.basename(src_path)}: +{c['added']} added, {c['skipped']} duplicates skipped, {c['replaced']} replaced by a more complete copy")
    merged["modillions"].sort(key=sort_key)
    out = os.path.join(out_dir, "project.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(merged, f, indent=2)
    print(f"wrote {out} with {len(merged['modillions'])} parts")
    return out


# ----------------------------------------------------------------------------- report
def parse_readings(txt):
    vals = [float(x) for x in re.findall(r"\d*\.?\d+", str(txt or ""))]
    return vals


def build(project_json, out_docx):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p = load_project(project_json) if not isinstance(project_json, dict) else project_json
    s = p["settings"]; name = s.get("element_name") or "Part"; gname = s.get("group_name") or "Area"; abbr = gname[0].upper()
    units = s.get("units") or "in"; dec = 2 if units == "mm" else 3
    photo_w, photo_h = float(s.get("photo_w_in") or 3), float(s.get("photo_h_in") or 2)
    max_px = int(s.get("report_px") or s.get("store_px") or 1600)
    fig = int(s.get("figure_start") or 1)
    types = {t["code"]: t for t in p.get("type_list", [])}
    defects = {d["label"]: d["text"] for d in p.get("defect_list", [])}
    tmp = tempfile.mkdtemp(prefix="member_")

    def tname(m): return (types.get(m.get("type")) or {}).get("name", m.get("type", ""))
    def sdesc(m, key):
        sf = m["surfaces"].get(key, {})
        if sf.get("desc"): return sf["desc"]
        for x in (types.get(m.get("type")) or {}).get("surfaces", []):
            if x["key"] == key: return x["desc"]
        return ""
    def checked(m):
        order = [x["key"] for x in (types.get(m.get("type")) or {}).get("surfaces", [])]
        ks = [k for k in m.get("checked", []) if k in m["surfaces"]]
        return sorted(ks, key=lambda k: order.index(k) if k in order else 99)

    # ---- helpers (same as the modillion tool) ----
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
    def row_h(row, inches, rule="atLeast"):
        trPr = row._tr.get_or_add_trPr(); h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(inches * 1440))); h.set(qn("w:hRule"), rule); trPr.append(h)
    def keep(row):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    def caption(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9):
        pp = cell.paragraphs[0]; pp.alignment = align; pp.paragraph_format.space_before = Pt(2); pp.paragraph_format.space_after = Pt(2)
        r = pp.add_run(text); r.font.size = Pt(size); return pp
    def prep(src):
        im = Image.open(src)
        try: orient = im.getexif().get(0x0112, 1)
        except Exception: orient = 1
        if im.format == "JPEG" and orient in (1, None) and max(im.size) <= max_px:
            return src, im.width, im.height
        im = ImageOps.exif_transpose(im)
        if im.mode not in ("RGB", "L"): im = im.convert("RGB")
        im.thumbnail((max_px, max_px)); out = os.path.join(tmp, f"img_{abs(hash(src)) & 0xFFFFFFFF:08x}.jpg"); im.save(out, "JPEG", quality=92)
        return out, im.width, im.height
    def fit(w, h, mw, mh):
        sc = min(mw / w, mh / h); return w * sc, h * sc
    def add_photo(cell, rel, mw, mh, align=WD_ALIGN_PARAGRAPH.CENTER, para=None):
        pp = para or cell.paragraphs[0]; pp.alignment = align; pp.paragraph_format.space_before = Pt(4); pp.paragraph_format.space_after = Pt(4)
        src = os.path.join(p["_photo_dir"], rel) if rel else ""
        if not rel or not os.path.isfile(src):
            r = pp.add_run("[photo missing]"); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xC0, 0, 0); return False
        path, w, h = prep(src); fw, fh = fit(w, h, mw, mh); pp.add_run().add_picture(path, width=Inches(fw), height=Inches(fh)); return True
    def page_field(paragraph):
        run = paragraph.add_run(); run.bold = True
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag: fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
            else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)

    doc = Document()
    sec = doc.sections[-1]; sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Inches(1))
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
    for lvl in (1, 2):
        st = doc.styles[f"Heading {lvl}"]; st.font.color.rgb = RGBColor(0x4E, 0x8A, 0x3E); st.font.name = "Calibri Light"; st.font.bold = False; st.font.size = Pt(16 if lvl == 1 else 13)
    footer = sec.footer; fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = fp._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4"); top.set(qn("w:space"), "1"); top.set(qn("w:color"), "D9D9D9"); pbdr.append(top); pPr.append(pbdr)
    page_field(fp); fp.add_run(" | ").bold = True; r = fp.add_run("P a g e"); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    footer.add_paragraph(s.get("report_id", "")).alignment = WD_ALIGN_PARAGRAPH.LEFT
    content_w, col_w = 6.5, 3.25
    mods = sorted(p["modillions"], key=sort_key)
    zones = []
    for m in mods:
        if m["zone"] not in zones: zones.append(m["zone"])

    # ---- summary table ----
    if s.get("include_summary", True):
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.add_run(f"Table 1: Summary of defects by {gname.lower()}").font.size = Pt(9)
        cw = [1.2, 1.2, 1.4, 2.7]
        st = doc.add_table(rows=1, cols=4); st.style = "Table Grid"; st.alignment = WD_TABLE_ALIGNMENT.CENTER; st.autofit = False
        for i, t in enumerate((gname, name, "Surface", "Defects")):
            c = st.rows[0].cells[i]; c.width = Inches(cw[i]); shade(c, "808080"); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = pp.add_run(t); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        row_h(st.rows[0], 0.4)
        for z in zones:
            row = st.add_row(); row_h(row, 0.32); c = row.cells[0].merge(row.cells[3]); shade(c, "BFBFBF")
            pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = pp.add_run(f"{gname} {z}"); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            any_ = False
            for m in [x for x in mods if x["zone"] == z]:
                for key in checked(m):
                    defs = m["surfaces"][key].get("defects") or []
                    if not defs: continue
                    any_ = True; row = st.add_row(); row_h(row, 0.4); keep(row)
                    vals = (f"{gname} {z}", str(m["number"]), f"{key} - {sdesc(m, key)}", " ".join(defects.get(d, d) for d in defs))
                    for i, v in enumerate(vals):
                        c = row.cells[i]; c.width = Inches(cw[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        pp = c.paragraphs[0]
                        if i in (0, 1): pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if i == 3: shade(c, "F2F2F2")
                        pp.add_run(v).font.size = Pt(11)
            if not any_:
                row = st.add_row(); c = row.cells[0].merge(row.cells[3]); pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = pp.add_run("No defects recorded"); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- areas ----
    first = True
    for zi, z in enumerate(zones, start=1):
        if not first: doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        zinfo = p.get("zones", {}).get(z, {})
        secno = str(s.get("section_number") or "").strip()
        doc.add_heading(f"{secno}.{zi}\t{gname} {z}:" if secno else f"{gname} {z}:", level=2)
        if (zinfo.get("description") or "").strip(): doc.add_paragraph(zinfo["description"].strip())
        zimg = os.path.join(p["_photo_dir"], zinfo.get("image", "")) if zinfo.get("image") else ""
        if zimg and os.path.isfile(zimg):
            path, w, h = prep(zimg); fw, fh = fit(w, h, content_w, 6.0)
            pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(path, width=Inches(fw), height=Inches(fh))
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.add_run(f"Figure {fig}: {(zinfo.get('image_caption') or '').strip() or f'{gname} {z} layout.'}").font.size = Pt(9); fig += 1
        for m in [x for x in mods if x["zone"] == z]:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            lbl = label(m, abbr)
            t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
            row = t.add_row(); row_h(row, 0.42); hc = row.cells[0].merge(row.cells[1]); shade(hc, "808080"); hc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pp = hc.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for txt, b, sz in ((f"{gname} {m['zone']}: ", False, 14), (f"{name} {m['number']}", True, 14), (f"  ({tname(m)})", False, 11)):
                r = pp.add_run(txt); r.bold = b; r.font.size = Pt(sz); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # overall | general notes
            row = t.add_row(); keep(row); row_h(row, photo_h + 0.15)
            c = row.cells[0]; c.width = Inches(col_w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            ok = bool(m["overall"]) and add_photo(c, m["overall"][0], photo_w, photo_h); ofig = 0
            if ok: ofig = fig; fig += 1
            c = row.cells[1]; c.width = Inches(col_w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            pp = c.paragraphs[0]; pp.paragraph_format.space_before = Pt(4)
            for k, line in enumerate((m.get("notes") or "").strip().splitlines() or [""]):
                if k: pp = c.add_paragraph()
                pp.add_run(line).font.size = Pt(11)
            row = t.add_row(); row_h(row, 0.22)
            caption(row.cells[0], f"Figure {ofig}: {name} {lbl} Overall" if ofig else f"{name} {lbl} Overall (no photo)"); caption(row.cells[1], "Notes")
            extras = m["overall"][1:]
            for i in range(0, len(extras), 2):
                row = t.add_row(); keep(row); crow = t.add_row(); row_h(crow, 0.22)
                for j, rel in enumerate(extras[i:i + 2]):
                    c = row.cells[j]; c.width = Inches(col_w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if add_photo(c, rel, col_w - 0.3, 1.7): caption(crow.cells[j], f"Figure {fig}: {name} {lbl} Overall ({i + j + 2})"); fig += 1
            # surfaces
            for key in checked(m):
                sf = m["surfaces"][key]; desc = sdesc(m, key)
                row = t.add_row(); keep(row); row_h(row, photo_h + 0.15)
                c = row.cells[0]; c.width = Inches(col_w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pf = 0
                if sf.get("photo") and add_photo(c, sf["photo"], photo_w, photo_h): pf = fig; fig += 1
                c = row.cells[1]; c.width = Inches(col_w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                vals = parse_readings(sf.get("readings"))
                pp = c.paragraphs[0]; pp.paragraph_format.space_before = Pt(4)
                if vals:
                    r = pp.add_run("Thickness: "); r.bold = True; r.font.size = Pt(11)
                    pp.add_run(", ".join(f"{v:.{dec}f}" for v in vals) + f" {units}").font.size = Pt(11)
                    p2 = c.add_paragraph(); r = p2.add_run(f"Min {min(vals):.{dec}f} {units}" + (f", avg {sum(vals)/len(vals):.{dec}f} {units}" if len(vals) > 1 else "")); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
                else:
                    r = pp.add_run("Thickness: not measured"); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
                for line in (sf.get("notes") or "").strip().splitlines():
                    c.add_paragraph().add_run(line).font.size = Pt(11)
                if sf.get("gauge"):
                    gp = c.add_paragraph()
                    if add_photo(c, sf["gauge"], 2.2, 1.5, WD_ALIGN_PARAGRAPH.LEFT, para=gp):
                        r = c.add_paragraph().add_run(f"Figure {fig}: Gauge reading, surface {key}"); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59); fig += 1
                row = t.add_row(); row_h(row, 0.22)
                caption(row.cells[0], f"Figure {pf}: {name} {lbl} - Surface {key}" + (f" ({desc})" if desc else "") if pf else f"{name} {lbl} - Surface {key} (no photo)")
                caption(row.cells[1], f"Surface {key}" + (f" - {desc}" if desc else ""))
    doc.save(out_docx)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"wrote {out_docx} ({len(mods)} parts)")


# ----------------------------------------------------------------------------- cli
def main(argv):
    a = argv[1:]
    if not a or a[0] in ("-h", "--help"):
        print(__doc__); return
    merged = None
    if "--merge" in a:
        i = a.index("--merge"); j = a.index("--build") if "--build" in a[i:] else len(a)
        out_dir, inputs = a[i + 1], a[i + 2:j]
        if not inputs: print("usage: --merge <output folder> export1.zip export2.zip ..."); sys.exit(2)
        merged = merge(out_dir, inputs); a = a[:i] + a[j:]
    if "--build" in a:
        i = a.index("--build"); rest = a[i + 1:]
        if merged and len(rest) == 1: build(merged, rest[0])
        elif len(rest) == 2: build(rest[0], rest[1])
        else: print("usage: --build project.json Report.docx   (or --merge ... --build Report.docx)"); sys.exit(2)


if __name__ == "__main__":
    main(sys.argv)
